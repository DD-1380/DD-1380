"""Checks scoring.py against the hand-graded cards in the USU/UCF spreadsheet.

    python tools/validate_scoring.py path/to/USU_UCF_Scored_DD1380.xlsx

Each column of the spreadsheet is one card, transcribed item by item, with
human 30-item, 11-item and critical totals. The transcription is turned into
pipeline field values, scored, and compared with the human scores.
"""
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scan-pipeline"))

from scoring import (  # noqa: E402
    AIRWAY, BREATHING, CIRCULATION, CRITICAL_11, EVAC, EVAC_BACK, FLUID_ROWS, GENDER,
    MECHANISMS, MED_ROWS, OTHER, TOURNIQUETS, VITALS, score_11, score_30,
)

# Spreadsheet row of each transcribed item.
ROWS = {
    "br": 2, "evac": 3, "name": 4, "last4": 5, "gender": 6, "date": 7, "time": 8, "service": 9,
    "unit": 10, "allergies": 11, "mechanism": 12, "diagram": 13,
    "ss_time": 14, "ss_pulse": 15, "ss_bp": 16, "ss_resp": 17, "ss_pulse_ox": 18, "ss_avpu": 19, "ss_pain": 20,
    "br_back": 22, "evac_back": 23, "circulation": 24, "airway": 25, "breathing": 26, "fluid": 27,
    "meds": 28, "other": 29, "notes": 30, "fr_name": 31, "fr_last4": 32,
}
TOTAL_30_ROW, ITEMS_11_FIRST_ROW, TOTAL_11_ROW, CRITICAL_ROW = 33, 35, 46, 47


def text(cell) -> str:
    return "" if cell is None else str(cell).strip()


EVAC_WORDS = ["URGENT", "PRIORITY", "ROUTINE"]


def evac_boxes(value: str, keys: list[str]) -> dict:
    return {k: word in value.upper() for k, word in zip(keys, EVAC_WORDS)}


# Turns one transcribed card into pipeline field values, with every field
# extracted (blank or checked/unchecked).
def card_values(column: dict) -> dict:
    values = {}
    for keys in (EVAC, EVAC_BACK, GENDER, MECHANISMS, CIRCULATION, AIRWAY, BREATHING, OTHER):
        values.update({k: (False if k.startswith("field_image_") else "") for k in keys})
    for row in FLUID_ROWS + MED_ROWS:
        values.update({k: "" for k in row})
    values.update({k: "" for pair in TOURNIQUETS for k in pair})

    values.update({
        "field_battle_roster": column["br"], "field_name": column["name"], "field_last_4": column["last4"],
        "field_date": column["date"], "field_time": column["time"], "field_service": column["service"],
        "field_unit": column["unit"], "field_allergies": column["allergies"],
        "field_battle_roster_back": column["br_back"],
        "field_tq_type": column["circulation"], "field_airway_type": column["airway"],
        "field_breathing_type": column["breathing"], "field_fluid_name_1": column["fluid"],
        "field_othermed_name_1": column["meds"], "field_hypothermia_type": column["other"],
        "field_notes": column["notes"], "field_first_responder_name": column["fr_name"],
        "field_first_responder_last_4": column["fr_last4"],
        "field_image_male": column["gender"].upper().startswith("M"),
        "field_image_female": column["gender"].upper().startswith("F"),
        "field_injury_body": column["diagram"].upper() == "YES",
    })
    values.update(evac_boxes(column["evac"], EVAC))
    values.update(evac_boxes(column["evac_back"], EVAC_BACK))

    # Mechanism: tick each named box; anything else is "Other" with the text.
    mechanism = column["mechanism"]
    for key in MECHANISMS:
        values[key] = re.search(rf"\b{key.split('_')[2].upper()}\b", mechanism.upper()) is not None
    named = any(values[k] for k in MECHANISMS)
    values["field_image_other"] = "OTHER" in mechanism.upper() or (bool(mechanism) and not named)
    values["field_other"] = mechanism if values["field_image_other"] else ""

    for name, keys in VITALS.items():
        values.update({k: "" for k in keys})
        values[keys[0]] = column[name]
    return values


def main(path: str) -> None:
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    names_11 = ["battle_roster_num", "evac_category", "patient_info", "mechanism_of_injury", "injuries",
                "signs_and_symptoms", "treatment_info", "meds", "other", "notes", "first_responder_identity"]
    matches = {"30-item": 0, "11-item": 0, "critical": 0}
    cards = 0

    for col in range(2, sheet.max_column + 1):
        card = text(sheet.cell(1, col).value)
        if not card:
            continue
        cards += 1
        column = {k: text(sheet.cell(r, col).value) for k, r in ROWS.items()}
        values = card_values(column)
        items_30 = score_30(values)
        items_11 = score_11(values, items_30)

        ours = {
            "30-item": sum(bool(v) for v in items_30.values()),
            "11-item": sum(bool(v) for v in items_11.values()),
            "critical": sum(bool(items_11[k]) for k in CRITICAL_11),
        }
        human = {
            "30-item": int(sheet.cell(TOTAL_30_ROW, col).value or 0),
            "11-item": int(sheet.cell(TOTAL_11_ROW, col).value or 0),
            "critical": int(sheet.cell(CRITICAL_ROW, col).value or 0),
        }
        status = []
        for k in ours:
            matches[k] += ours[k] == human[k]
            status.append(f"{k} {ours[k]}/{human[k]}" + ("" if ours[k] == human[k] else " MISMATCH"))
        print(f"{card:<16} " + " | ".join(status))

        for i, name in enumerate(names_11):
            human_item = int(sheet.cell(ITEMS_11_FIRST_ROW + i, col).value or 0)
            if bool(items_11[name]) != bool(human_item):
                print(f"    11-item '{name}': ours {int(bool(items_11[name]))}, human {human_item}")

    print("\nagreement: " + ", ".join(f"{k} {n}/{cards}" for k, n in matches.items()))


if __name__ == "__main__":
    main(sys.argv[1])
